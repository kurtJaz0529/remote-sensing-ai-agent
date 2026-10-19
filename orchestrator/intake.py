"""Read-only file intake with explicit extraction gaps and requirement locations."""
from pathlib import Path
from contracts import VERSION, sha256, validate, within

def inventory(root, dataset_id):
    root = Path(root).resolve()
    entries = []
    for path in sorted(root.rglob("*")):
        if path.is_file():
            relative = path.relative_to(root).as_posix()
            resolved = within(root, relative)
            entries.append(dict(relative_path=relative, sha256=sha256(resolved), size=resolved.stat().st_size))
    return validate("input_manifest", dict(schema_revision=VERSION, dataset_id=dataset_id, origin="user_provided", files=entries))

def check_inputs(root, manifest):
    validate("input_manifest", manifest)
    for entry in manifest["files"]:
        path = within(root, entry["relative_path"])
        if not path.is_file() or path.stat().st_size != entry["size"] or sha256(path) != entry["sha256"]:
            raise ValueError("input changed: " + entry["relative_path"])

def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    blocks, gaps = [], []
    if suffix in (".txt", ".md"):
        for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if line.strip():
                blocks.append(dict(source_locator=f"{path.name}:line {number}", text=line.strip(), confidence=1.0))
    elif suffix == ".pdf":
        from pypdf import PdfReader
        for number, page in enumerate(PdfReader(path).pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                blocks.append(dict(source_locator=f"{path.name}:page {number}", text=text, confidence=0.7))
            else:
                gaps.append(f"{path.name}:page {number}: no extractable text; OCR/manual review required")
        gaps.append("PDF figures, equations, layout and attachments require review; text extraction is not full interpretation")
    elif suffix == ".docx":
        from docx import Document
        doc = Document(path)
        for number, para in enumerate(doc.paragraphs, 1):
            if para.text.strip():
                blocks.append(dict(source_locator=f"{path.name}:paragraph {number}", text=para.text, confidence=1.0))
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                blocks.append(dict(source_locator=f"{path.name}:table {ti}:row {ri}", text=" | ".join(c.text for c in row.cells), confidence=1.0))
        gaps.append("DOCX images, equations, headers and embedded attachments require manual review")
    else:
        raise ValueError("unsupported requirement format; use UTF-8 TXT/MD, DOCX or text PDF")
    return dict(document=path.name, sha256=sha256(path), blocks=blocks, gaps=gaps,
                interpretation_status="requires_review", conflicts=[], personal_requirements=[])

def validate_plan(plan):
    validate("task", plan)
    steps = {s["step_id"]: s for s in plan["method_plan"]}
    if len(steps) != len(plan["method_plan"]):
        raise ValueError("duplicate step IDs")
    if len({r["requirement_id"] for r in plan["requirements"]}) != len(plan["requirements"]):
        raise ValueError("duplicate requirement IDs")
    seen, active = set(), set()
    def visit(sid):
        if sid in active:
            raise ValueError("cyclic step dependency")
        if sid in seen:
            return
        if sid not in steps:
            raise ValueError("unknown dependency: " + sid)
        active.add(sid)
        for dep in steps[sid]["depends_on"]:
            visit(dep)
        active.remove(sid)
        seen.add(sid)
    for sid, step in steps.items():
        visit(sid)
        # Validate identifiers/paths before they become filenames.
        if not sid.replace("-", "").replace("_", "").isalnum():
            raise ValueError("nonportable step ID")
        if step["output_dir"] != "artifacts/" + sid:
            raise ValueError("output_dir must equal artifacts/<step_id>")
        if not set(step["required_kinds"]) <= {"data","project","reproducible_package","screenshot","docx","pdf","log","check"}:
            raise ValueError("unknown artifact kind")
    required=set(plan["required_deliverables"])
    for requirement in plan["requirements"]:
        if not requirement["required_artifacts"] or not requirement["acceptance_checks"]:
            raise ValueError("each requirement needs explicit deliverables and checks")
        if not set(requirement["required_artifacts"]) <= required:
            raise ValueError("requirement artifacts missing from deliverable gate")
    if plan["open_questions"] or plan["conflicts"]:
        raise ValueError("unresolved questions/conflicts prevent approval")
    return plan
