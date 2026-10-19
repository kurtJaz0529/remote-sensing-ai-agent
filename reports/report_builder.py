"""Evidence-bound experiment report verifier and DOCX/PDF builder (T23)."""
from __future__ import annotations

import re
import math
from pathlib import Path
from typing import Any, Mapping, Sequence

from docx import Document
from pypdf import PdfReader

from contracts import sha256, within

FORBIDDEN_HEADLESS_GUI_PHRASES = (
    "点击菜单栏",
    "在弹出的对话框中勾选",
    "鼠标右键选择",
    "手动拖拽窗口截图",
    "在 Toolbox 界面双击打开",
    "clicked the menu bar",
    "in the popup dialog",
)

FORBIDDEN_OVERCLAIM_ACCURACY_PHRASES = (
    "独立样本真实泛化精度",
    "独立测试集泛化精度",
    "全图无偏面积精度",
    "independent generalization accuracy",
    "unbiased map area accuracy",
)

FORBIDDEN_METRIC_AREA_REGEXES = (
    r"\b\d+(?:\.\d+)?\s*(?:km\^?2|hm\^?2|ha)\b",
    r"\b\d+(?:\.\d+)?\s*(?:公顷|平方千米|平方公里)\b",
)


def audit_report_table(
    table_spec: Mapping[str, Any],
    manifest_metrics: Mapping[str, Mapping[str, Any]],
    tolerance: float = 1e-2,
) -> dict[str, Any]:
    """Verify report comparison table column headers and row values against manifest bindings.

    Prevents historical ENVI Ex8 Table 6 errors where ML_Clump and ML_Sieve columns or
    built_up and water class columns were swapped.
    """
    issues: list[dict[str, str]] = []
    columns: Sequence[Mapping[str, str]] = table_spec.get("columns", [])
    rows: Sequence[Mapping[str, Any]] = table_spec.get("rows", [])
    column_ids = [str(c.get("column_id", "")) for c in columns]
    row_ids = [str(r.get("row_id", "")) for r in rows]
    if not column_ids or not row_ids or "" in column_ids + row_ids or len(set(column_ids)) != len(column_ids) or len(set(row_ids)) != len(row_ids):
        issues.append({"check_id": "CHK-REPORT-TABLE-COL-001", "severity": "error", "message": "Table rows and columns must be nonempty and uniquely bound"})

    # Check column header label vs bound key consistency
    for col in columns:
        col_id = str(col.get("column_id", ""))
        bound_id = str(col.get("bound_id", col_id))
        header_text = str(col.get("header", ""))
        if col_id and bound_id and col_id != bound_id:
            issues.append(
                {
                    "check_id": "CHK-REPORT-TABLE-COL-001",
                    "severity": "error",
                    "message": (
                        f"Column binding mismatch: column_id={col_id!r} (header={header_text!r}) "
                        f"is bound to {bound_id!r}."
                    ),
                }
            )

    for row in rows:
        row_id = str(row.get("row_id", ""))
        bound_row_id = str(row.get("bound_row_id", row_id))
        if row_id != bound_row_id:
            issues.append(
                {
                    "check_id": "CHK-REPORT-TABLE-COL-001",
                    "severity": "error",
                    "message": f"Row binding mismatch: row_id={row_id!r} bound to {bound_row_id!r}.",
                }
            )
            continue

        cells: Mapping[str, Any] = row.get("cells", {})
        if row_id not in manifest_metrics or set(cells) != set(column_ids) or any(c not in manifest_metrics.get(row_id, {}) for c in column_ids):
            issues.append({"check_id": "CHK-REPORT-TABLE-COL-001", "severity": "error", "message": "Table contains missing or unknown metric bindings"})
        if row_id in manifest_metrics:
            expected_row = manifest_metrics[row_id]
            for col in columns:
                col_id = str(col["column_id"])
                if col_id in cells and col_id in expected_row:
                    rep_val = float(cells[col_id])
                    exp_val = float(expected_row[col_id])
                    if not math.isfinite(rep_val) or not math.isfinite(exp_val):
                        issues.append({"check_id": "CHK-REPORT-TABLE-COL-001", "severity": "error", "message": "Table metrics must be finite"})
                        continue
                    if abs(rep_val - exp_val) > tolerance:
                        # Check if it matches another column (column swap!)
                        swapped_with = [
                            other_col
                            for other_col, other_val in expected_row.items()
                            if other_col != col_id and abs(rep_val - float(other_val)) <= tolerance
                        ]
                        swap_note = (
                            f" (value matches column {swapped_with[0]!r} -> column swap detected!)"
                            if swapped_with
                            else ""
                        )
                        issues.append(
                            {
                                "check_id": "CHK-REPORT-TABLE-COL-001",
                                "severity": "error",
                                "message": (
                                    f"Table value mismatch at row={row_id!r}, col={col_id!r}: "
                                    f"reported {rep_val}, manifest expected {exp_val}{swap_note}."
                                ),
                            }
                        )

    return {"passed": len(issues) == 0, "issues": issues}


def audit_report_payload(
    report_payload: Mapping[str, Any],
    manifest: Mapping[str, Any],
    artifact_root: str | Path | None = None,
) -> dict[str, Any]:
    """Audit report narrative, tables, figures, execution mode honesty, and evaluation_kind claims."""
    issues: list[dict[str, str]] = []
    if report_payload.get("delivery_kind", "formal_experiment") != "technical_preview":
        required = manifest.get("required_artifact_ids", [])
        present = {a.get("artifact_id") for a in manifest.get("artifacts", [])}
        ready = all(manifest.get(k) is True for k in ("formal_report_ready", "evidence_gate_complete", "sample_quality_passed"))
        ready = ready and bool(required) and set(required) <= present and artifact_root is not None
        ready = ready and all(re.fullmatch(r"[0-9a-f]{64}", str(manifest.get(k, ""))) for k in ("plan_hash", "input_manifest_hash", "template_sha256"))
        if not ready:
            issues.append({"check_id": "CHK-REPORT-READY-001", "severity": "error", "message": "Formal experiment report requires a frozen plan/template, passed sample quality and complete required evidence"})
    execution_mode = str(manifest.get("execution_mode", "headless_api"))
    evaluation_kind = str(manifest.get("evaluation_kind", "reference_agreement"))
    has_crs = bool(manifest.get("has_crs", False))
    has_verified_scale = bool(manifest.get("has_verified_scale_factor", False))

    narrative_sections: Mapping[str, str] = report_payload.get("sections", {})
    full_text = "\n".join(str(v) for v in narrative_sections.values())

    # 1. Execution mode honesty check (API/batch vs GUI)
    if execution_mode in ("headless_api", "batch"):
        for phrase in FORBIDDEN_HEADLESS_GUI_PHRASES:
            if phrase.lower() in full_text.lower():
                issues.append(
                    {
                        "check_id": "CHK-REPORT-HONESTY-001",
                        "severity": "error",
                        "message": (
                            f"Forbidden GUI interaction phrase {phrase!r} found in report while "
                            f"execution_mode={execution_mode!r}."
                        ),
                    }
                )

    # 2. Evaluation kind honesty check
    if evaluation_kind != "buffered_independent_holdout":
        for phrase in FORBIDDEN_OVERCLAIM_ACCURACY_PHRASES:
            if phrase.lower() in full_text.lower():
                issues.append(
                    {
                        "check_id": "CHK-EVAL-KIND-001",
                        "severity": "error",
                        "message": (
                            f"Forbidden overclaim phrase {phrase!r} found while "
                            f"evaluation_kind={evaluation_kind!r}."
                        ),
                    }
                )

    # 3. Missing CRS metric area check
    if not has_crs:
        for pattern in FORBIDDEN_METRIC_AREA_REGEXES:
            match = re.search(pattern, full_text, flags=re.IGNORECASE)
            if match:
                issues.append(
                    {
                        "check_id": "CHK-AUDIT-CRS-001",
                        "severity": "error",
                        "message": (
                            f"Metric area expression {match.group(0)!r} found in report despite "
                            f"has_crs=False."
                        ),
                    }
                )

    # 4. Unverified reflectance scale check
    if not has_verified_scale:
        if "已定标地表反射率" in full_text or "calibrated surface reflectance (0..1)" in full_text.lower():
            issues.append(
                {
                    "check_id": "CHK-AUDIT-SCALE-001",
                    "severity": "error",
                    "message": "Report claims calibrated 0..1 surface reflectance without verified scale factor.",
                }
            )

    # 5. K-T third component naming check
    spectral_text = str(narrative_sections.get("spectral_transform", ""))
    if ("缨帽" in spectral_text or "K-T" in spectral_text or "Tasseled Cap" in spectral_text):
        if "Third" not in spectral_text:
            issues.append(
                {
                    "check_id": "CHK-SPEC-KT-THIRD-001",
                    "severity": "error",
                    "message": (
                        "K-T transform section must explicitly designate Landsat 5 TM 3rd component "
                        "as 'Third (Wetness-related)'."
                    ),
                }
            )

    # 6. Tables audit
    manifest_tables = manifest.get("table_metrics", {})
    for tbl in report_payload.get("tables", []):
        tbl_id = str(tbl.get("table_id", "default"))
        expected_metrics = manifest_tables.get(tbl_id, {})
        tbl_res = audit_report_table(tbl, expected_metrics)
        issues.extend(tbl_res["issues"])

    # 7. Figures and screenshot SHA256 audit
    manifest_artifacts = {
        str(a["artifact_id"]): a for a in manifest.get("artifacts", []) if "artifact_id" in a
    }
    for fig in report_payload.get("figures", []):
        art_id = str(fig.get("artifact_id", ""))
        if art_id not in manifest_artifacts:
            issues.append(
                {
                    "check_id": "CHK-EVIDENCE-MANIFEST-001",
                    "severity": "error",
                    "message": f"Figure references unknown artifact_id={art_id!r} not in frozen manifest.",
                }
            )
            continue
        expected_art = manifest_artifacts[art_id]
        if not fig.get("sha256") or not expected_art.get("sha256") or artifact_root is None:
            issues.append({"check_id": "CHK-SCREENSHOT-ID-001", "severity": "error", "message": "Figures require original bytes, a hash and a readable artifact root"})
        if fig.get("sha256") and fig["sha256"] != expected_art.get("sha256"):
            issues.append(
                {
                    "check_id": "CHK-SCREENSHOT-ID-001",
                    "severity": "error",
                    "message": (
                        f"Figure SHA256 mismatch for artifact_id={art_id!r}: "
                        f"figure has {fig.get('sha256')}, manifest has {expected_art.get('sha256')}."
                    ),
                }
            )
        if artifact_root is not None and expected_art.get("relative_path"):
            try:
                file_path = within(artifact_root, str(expected_art["relative_path"]))
            except ValueError:
                issues.append({"check_id": "CHK-SCREENSHOT-ID-001", "severity": "error", "message": "Figure path escapes the artifact root"})
                continue
            if not file_path.exists():
                issues.append(
                    {
                        "check_id": "CHK-SCREENSHOT-ID-001",
                        "severity": "error",
                        "message": f"Referenced figure file missing on disk: {file_path}",
                    }
                )
            elif expected_art.get("sha256") and sha256(file_path) != expected_art["sha256"]:
                issues.append({"check_id": "CHK-SCREENSHOT-ID-001", "severity": "error", "message": f"Disk SHA256 mismatch for figure file: {file_path}"})

    if report_payload.get("delivery_kind", "formal_experiment") != "technical_preview" and artifact_root is not None:
        for art_id in manifest.get("required_artifact_ids", []):
            art = manifest_artifacts.get(art_id, {})
            try:
                path = within(artifact_root, str(art.get("relative_path", "")))
                valid = path.is_file() and bool(art.get("sha256")) and sha256(path) == art["sha256"]
            except ValueError:
                valid = False
            if not valid:
                issues.append({"check_id": "CHK-REPORT-READY-001", "severity": "error", "message": "Required evidence bytes are absent or differ from the frozen manifest"})

    return {
        "passed": len(issues) == 0,
        "run_id": manifest.get("run_id"),
        "execution_mode": execution_mode,
        "evaluation_kind": evaluation_kind,
        "issues": issues,
    }


def build_experiment_docx(
    report_payload: Mapping[str, Any],
    manifest: Mapping[str, Any],
    output_docx_path: str | Path,
    artifact_root: str | Path | None = None,
) -> dict[str, Any]:
    """Validate report payload against frozen manifest and generate a structured .docx report."""
    if report_payload.get("delivery_kind") != "technical_preview":
        raise ValueError("Formal template-preserving report generation is not verified by this candidate builder")
    audit = audit_report_payload(report_payload, manifest, artifact_root=artifact_root)
    if not audit["passed"]:
        raise ValueError(f"Report payload failed verification checks: {audit['issues']}")

    doc = Document()
    title = str(report_payload.get("title", "遥感图像分析与分类实验报告"))
    doc.add_heading(title, level=0)
    doc.add_paragraph("Technical preview / 技术预览：不代表正式实验报告或独立验收。")

    meta_para = doc.add_paragraph()
    meta_para.add_run(f"Run ID: {manifest.get('run_id', 'N/A')} | ")
    meta_para.add_run(f"Execution Mode: {manifest.get('execution_mode', 'headless_api')} | ")
    meta_para.add_run(f"Evaluation Kind: {manifest.get('evaluation_kind', 'reference_agreement')} | ")
    meta_para.add_run(f"Review Status: {manifest.get('review_status', 'pending_review')}")

    section_titles = [
        ("task_and_env", "1. 任务契约与本地软件环境探针"),
        ("dataset_audit", "2. 数据集元数据、空间参考与样本缓冲审计"),
        ("spectral_transform", "3. 光谱指数、PCA 与 K-T 变换分析"),
        ("classification", "4. 非监督与监督分类方法及参数边界"),
        ("postprocessing", "5. 分类后处理多分支对比与水体保持分析"),
        ("accuracy_evaluation", "6. 分类精度验证与混淆矩阵分析"),
        ("evidence_manifest", "7. 冻结运行证据索引与复核状态"),
    ]
    sections: Mapping[str, str] = report_payload.get("sections", {})
    for key, heading in section_titles:
        if key in sections:
            doc.add_heading(heading, level=1)
            doc.add_paragraph(str(sections[key]))

    for tbl in report_payload.get("tables", []):
        doc.add_heading(str(tbl.get("title", tbl.get("table_id", "数据对比表"))), level=2)
        cols = tbl.get("columns", [])
        rows = tbl.get("rows", [])
        if cols:
            word_tbl = doc.add_table(rows=1, cols=len(cols) + 1)
            hdr_cells = word_tbl.rows[0].cells
            hdr_cells[0].text = "类别 / 分支"
            for idx, c in enumerate(cols):
                hdr_cells[idx + 1].text = str(c.get("header", c.get("column_id", "")))
            for r in rows:
                row_cells = word_tbl.add_row().cells
                row_cells[0].text = str(r.get("label", r.get("row_id", "")))
                cell_map = r.get("cells", {})
                for idx, c in enumerate(cols):
                    cid = str(c["column_id"])
                    row_cells[idx + 1].text = str(cell_map.get(cid, ""))

    artifacts = {a["artifact_id"]: a for a in manifest.get("artifacts", [])}
    for fig in report_payload.get("figures", []):
        original = within(artifact_root, artifacts[fig["artifact_id"]]["relative_path"])
        doc.add_picture(str(original))
        doc.add_paragraph(str(fig.get("caption", fig["artifact_id"])))

    out_path = Path(output_docx_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {
        "passed": True,
        "docx_path": str(out_path),
        "docx_sha256": sha256(out_path),
        "size_bytes": out_path.stat().st_size,
    }


def verify_pdf_file(pdf_path: str | Path) -> dict[str, Any]:
    """Verify that a generated PDF file has a valid %PDF- header and at least one readable page."""
    p = Path(pdf_path)
    if not p.exists():
        return {"passed": False, "reason": "file_missing"}
    raw = p.read_bytes()
    if not raw.startswith(b"%PDF-"):
        return {"passed": False, "reason": "invalid_pdf_magic_header"}
    reader = PdfReader(str(p))
    page_count = len(reader.pages)
    return {
        "passed": page_count >= 1,
        "page_count": page_count,
        "pdf_sha256": sha256(p),
        "size_bytes": len(raw),
    }
