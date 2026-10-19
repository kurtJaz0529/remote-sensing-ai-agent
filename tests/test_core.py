import copy
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch
import httpx
from jsonschema import ValidationError, Draft202012Validator
from contracts import VERSION, canonical, fingerprint, sha256, schema, validate, within
from orchestrator.engine import Engine
from orchestrator.examples import mock_plan, mock_tool
from orchestrator.intake import inventory, extract
from orchestrator.mock_worker import MockWorker
from orchestrator.providers import Provider, ProviderError, strict_json
from orchestrator.worker import JsonWorker
from retrieval.index import KnowledgeIndex, compatible
from retrieval.experience import ExperienceRegistry

class RuntimeCase(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.base=Path(self.tmp.name); self.inputs=self.base/"inputs"; self.inputs.mkdir()
        (self.inputs/"sample.txt").write_text("user-provided test fixture",encoding="utf-8")
        self.manifest=inventory(self.inputs,"DATA-1")
        self.plan=mock_plan(self.manifest)
        self.engine=Engine(self.base/"runtime"); self.addCleanup(self.engine.store.close)
        self.worker=MockWorker(self.base/"runtime"); self.addCleanup(self.worker.close)

    def prepare(self, plan=None):
        plan=plan or self.plan
        self.engine.save_plan(plan,self.manifest,self.inputs)
        approved=self.engine.approve(plan["task_id"],plan["revision"],"test-user")
        return self.engine.create_run(plan["task_id"],plan["revision"],approved,"mock")

    def test_no_approval_no_execution(self):
        self.engine.save_plan(self.plan,self.manifest,self.inputs)
        with self.assertRaises(ValueError): self.engine.create_run("TASK-DEMO",1,"not-an-approval","mock")
        self.assertEqual(self.worker.submissions,0)

    def test_approval_bound_to_revision(self):
        self.engine.save_plan(self.plan,self.manifest,self.inputs)
        approval=self.engine.approve("TASK-DEMO",1,"test-user")
        p=copy.deepcopy(self.plan); p["revision"]=2
        self.engine.save_plan(p,self.manifest,self.inputs)
        with self.assertRaises(ValueError): self.engine.create_run("TASK-DEMO",2,approval,"mock")

    def test_revision_immutable(self):
        self.engine.save_plan(self.plan,self.manifest,self.inputs)
        self.plan["objective"]="changed"
        with self.assertRaises(ValueError): self.engine.save_plan(self.plan,self.manifest,self.inputs)

    def test_conflicts_and_unknowns_block(self):
        for field in ("conflicts","open_questions"):
            with self.subTest(field=field):
                p=copy.deepcopy(self.plan); p["task_id"]="TASK-"+field; p[field]=["unresolved"]
                self.engine.save_plan(p,self.manifest,self.inputs)
                with self.assertRaises(ValueError): self.engine.approve(p["task_id"],1,"test-user")

    def test_cycle_blocks(self):
        self.plan["method_plan"][0]["depends_on"]=["S1"]
        self.engine.save_plan(self.plan,self.manifest,self.inputs)
        with self.assertRaises(ValueError): self.engine.approve("TASK-DEMO",1,"test-user")

    def test_unknown_dependency_blocks(self):
        self.plan["method_plan"][0]["depends_on"]=["missing"]
        self.engine.save_plan(self.plan,self.manifest,self.inputs)
        with self.assertRaises(ValueError): self.engine.approve("TASK-DEMO",1,"test-user")

    def test_original_change_blocks_preflight(self):
        run=self.prepare(); (self.inputs/"sample.txt").write_text("changed")
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Blocked")
        self.assertEqual(self.worker.submissions,0)

    def test_invalid_parameters_never_dispatched(self):
        self.plan["method_plan"][0]["parameters"]["extra"]="unapproved"
        with self.assertRaises(ValidationError): self.prepare()
        self.assertEqual(self.worker.submissions,0)

    def test_real_mode_needs_local_capabilities(self):
        self.engine.save_plan(self.plan,self.manifest,self.inputs)
        approval=self.engine.approve("TASK-DEMO",1,"test-user")
        with self.assertRaises(ValueError): self.engine.create_run("TASK-DEMO",1,approval,"real")

    def test_successful_mock_has_evidence_and_preserves_input(self):
        run=self.prepare(); status=self.engine.execute(run,self.worker)
        self.assertEqual(status["state"],"Completed"); self.assertEqual(status["validation_scope"],"mock_only")
        self.assertEqual(inventory(self.inputs,"DATA-1"),self.manifest)
        root=Path(self.engine.export(run)); manifest=json.loads((root/"manifest.json").read_text())
        for entry in manifest["files"]:
            self.assertEqual(sha256(root/entry["relative_path"]),entry["sha256"])
        for line in (root/"events/events.jsonl").read_text().splitlines(): validate("event",json.loads(line))

    def test_duplicate_resume_does_not_recompute(self):
        run=self.prepare(); self.engine.execute(run,self.worker); self.engine.execute(run,self.worker)
        self.assertEqual(self.worker.submissions,1)

    def test_duplicate_run_creation_reuses_same_approval(self):
        run=self.prepare(); approval=self.engine.status(run)["approval"]
        self.assertEqual(self.engine.create_run("TASK-DEMO",1,approval,"mock"),run)

    def test_handoff_bundle_integrity_and_mock_publication_refusal(self):
        from orchestrator.evidence import audit_bundle, load_reviewed_evidence
        run=self.prepare(); self.engine.execute(run,self.worker)
        root=Path(self.engine.export(run))
        proof=audit_bundle(root); self.assertEqual(proof["mode"],"mock")
        with self.assertRaises(ValueError): load_reviewed_evidence([root])
        (root/"artifacts/S1/mock.txt").write_text("tampered")
        with self.assertRaises(ValueError): audit_bundle(root)

    def test_lost_reply_reconciles_without_duplicate_write(self):
        run=self.prepare(); self.worker.scenario="lost_reply"
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Blocked")
        self.worker.scenario="success"
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Completed")
        self.assertEqual(self.worker.submissions,1)

    def test_restart_uses_durable_state(self):
        run=self.prepare(); self.worker.scenario="lost_reply"; self.engine.execute(run,self.worker)
        other=Engine(self.base/"runtime"); self.addCleanup(other.store.close)
        self.worker.scenario="success"
        self.assertEqual(other.execute(run,self.worker)["state"],"Completed")
        self.assertEqual(self.worker.submissions,1)

    def test_unknown_job_never_resubmitted(self):
        run=self.prepare()
        class LostBeforeSubmit:
            calls=0
            def submit(w,job): w.calls+=1; raise TimeoutError()
            def status(w,job): raise ValueError("unknown")
        worker=LostBeforeSubmit()
        self.engine.execute(run,worker); self.engine.execute(run,worker)
        self.assertEqual(worker.calls,1); self.assertEqual(self.engine.status(run)["state"],"Blocked")

    def test_no_evidence_no_completion(self):
        run=self.prepare(); self.worker.scenario="no_evidence"
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Blocked")

    def test_missing_required_screenshot_blocks(self):
        self.plan["required_deliverables"].append("screenshot")
        run=self.prepare()
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Blocked")

    def test_unfulfilled_acceptance_check_blocks_completion(self):
        self.plan["requirements"][0]["acceptance_checks"].append("independent-scientific-check")
        run=self.prepare()
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Blocked")

    def test_worker_failure_terminal_without_retry(self):
        run=self.prepare(); self.worker.scenario="failed"
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Failed")
        self.engine.execute(run,self.worker); self.assertEqual(self.worker.submissions,1)

    def test_budget_blocks_second_step(self):
        self.plan["budget"]["max_tool_calls"]=1
        s=copy.deepcopy(self.plan["method_plan"][0]); s.update(step_id="S2",depends_on=["S1"],output_dir="artifacts/S2")
        self.plan["method_plan"].append(s)
        run=self.prepare(); self.assertEqual(self.engine.execute(run,self.worker)["state"],"Blocked")
        self.assertEqual(self.worker.submissions,1)

    def test_running_job_cancel_retains_evidence(self):
        run=self.prepare(); self.worker.scenario="running"
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Paused")
        self.assertEqual(self.engine.cancel(run,self.worker)["state"],"Cancelled")
        self.assertTrue(list((self.base/"runtime/runs"/run/"artifacts").rglob("mock.txt")))

    def test_async_poll_then_success(self):
        run=self.prepare(); self.worker.scenario="running"; self.engine.execute(run,self.worker)
        self.worker.scenario="success"
        self.assertEqual(self.engine.execute(run,self.worker)["state"],"Completed")
        self.assertEqual(self.worker.submissions,1)

    def test_result_identity_parameter_and_hash_attacks(self):
        original=self.worker.submit
        for index,mutator in enumerate([
            lambda r:r.update(run_id="old-run"),
            lambda r:r.update(implementation_kind="native"),
            lambda r:r.update(effective_parameters={"message":"changed"}),
            lambda r:r["outputs"][0].update(sha256="0"*64),
            lambda r:r["outputs"][0].update(relative_path="../outside.txt"),
            lambda r:r["outputs"][0].update(verified=False),
            lambda r:r.update(checks=[]),
            lambda r:r["outputs"][0].update(kind="screenshot"),
            lambda r:r["outputs"][0].update(kind="pdf"),
            lambda r:r["outputs"][0].update(kind="docx"),
        ]):
            with self.subTest(index=index):
                p=copy.deepcopy(self.plan); p["task_id"]="ATTACK-"+str(index)
                run=self.prepare(p)
                def altered(job):
                    result=original(job); mutator(result); return result
                self.worker.submit=altered
                self.assertEqual(self.engine.execute(run,self.worker)["state"],"Blocked")
        self.worker.submit=original

    def test_subprocess_worker_protocol(self):
        run=self.prepare()
        worker=JsonWorker([sys.executable,"-m","orchestrator.mock_worker","--root",str(self.base/"runtime")])
        self.assertEqual(self.engine.execute(run,worker)["state"],"Completed")

class ProviderCase(unittest.TestCase):
    def setUp(self):
        self.config=dict(schema_revision=VERSION,provider="openai",base_url="https://api.openai.com/v1",model="configured-test-model",api_key_env="RSAGENT_TEST_KEY",timeout_seconds=5,context_tokens=10000,tool_calling="verified",vision="unknown",extra_body={},pricing=None)
        self.env=patch.dict(os.environ,{"RSAGENT_TEST_KEY":"test-secret"}); self.env.start(); self.addCleanup(self.env.stop)
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        from orchestrator.store import Store
        self.store=Store(self.tmp.name); self.addCleanup(self.store.close)

    def provider(self, handler):
        provider=Provider(self.config,self.store,httpx.MockTransport(handler)); self.addCleanup(provider.close); return provider

    def response(self,arguments='{"message":"ok"}',name="mock__echo",code=200):
        return httpx.Response(code,json={"model":"observed-model","choices":[{"finish_reason":"tool_calls","message":{"content":None,"tool_calls":[{"id":"call-1","type":"function","function":{"name":name,"arguments":arguments}}]}}],"usage":{"prompt_tokens":10,"completion_tokens":20,"total_tokens":30}})

    def test_openai_and_deepseek_wire_and_usage(self):
        for provider_name in ("openai","deepseek"):
            with self.subTest(provider=provider_name):
                self.config["provider"]=provider_name
                def handler(request):
                    body=json.loads(request.content)
                    self.assertIn("max_completion_tokens" if provider_name=="openai" else "max_tokens",body)
                    return self.response()
                result=self.provider(handler).complete([{"role":"user","content":"test"}],[mock_tool()])
                self.assertEqual(result["tool_proposals"][0]["arguments"],{"message":"ok"})
                self.assertFalse(result["execution_authorized"])
        log=" ".join(r[0] for r in self.store.db.execute("SELECT body FROM usage"))
        self.assertIn("observed-model",log); self.assertNotIn("test-secret",log)

    def test_invalid_tool_arguments_never_released(self):
        for args in ('{"message":3}','{"message":"ok","extra":1}','{"message":"x","message":"y"}','{"message":NaN}','{broken'):
            with self.subTest(args=args), self.assertRaises(ProviderError):
                self.provider(lambda r:self.response(args)).complete([{"role":"user","content":"test"}],[mock_tool()])

    def test_unknown_tool_never_released(self):
        with self.assertRaises(ProviderError):
            self.provider(lambda r:self.response(name="shell_exec")).complete([{"role":"user","content":"test"}],[mock_tool()])

    def test_capability_guards(self):
        self.config["tool_calling"]="unknown"
        with self.assertRaises(ProviderError):
            self.provider(lambda r:self.response()).complete([{"role":"user","content":"test"}],[mock_tool()])
        with self.assertRaises(ProviderError):
            self.provider(lambda r:self.response()).complete([{"role":"user","content":[{"type":"image_url"}]}])

    def test_http_errors_are_classified_and_not_retried(self):
        for code,expected in ((401,"authentication"),(403,"authentication"),(429,"rate_limit"),(500,"server"),(400,"request")):
            with self.subTest(code=code):
                calls=[]
                def handler(request): calls.append(1); return self.response(code=code)
                with self.assertRaises(ProviderError) as raised:
                    self.provider(handler).complete([{"role":"user","content":"test"}])
                self.assertEqual(raised.exception.kind,expected); self.assertEqual(len(calls),1)

    def test_timeout_records_unknown_cost(self):
        def handler(request): raise httpx.ReadTimeout("test-secret")
        with self.assertRaises(ProviderError) as raised:
            self.provider(handler).complete([{"role":"user","content":"test"}])
        self.assertEqual(raised.exception.kind,"timeout")
        self.assertNotIn("test-secret",str(raised.exception))
        record=json.loads(self.store.db.execute("SELECT body FROM usage").fetchone()[0])
        self.assertIsNone(record["cost"])

    def test_missing_key_is_blocked_before_network(self):
        with patch.dict(os.environ,{},clear=True), self.assertRaises(ProviderError):
            self.provider(lambda r:self.fail("must not call network")).complete([{"role":"user","content":"test"}])

    def test_unknown_pricing_blocks_bounded_request(self):
        with self.assertRaises(ProviderError):
            self.provider(lambda r:self.fail("must not call network")).complete([{"role":"user","content":"test"}],max_cost=1)

    def test_persistent_budget_blocks_after_prior_cost_or_uncertainty(self):
        self.config["pricing"]={"input_per_million":1,"output_per_million":1,"source_url":"https://example.invalid/pricing","currency":"USD"}
        provider=self.provider(lambda r:self.response())
        # Prior spend includes previous provider/model calls under the same scope.
        self.store.db.execute("INSERT INTO usage(provider,model,timestamp,body) VALUES(?,?,?,?)",("deepseek","old","test",canonical({"budget_scope":"RUN-budget","cost":0.99,"currency":"USD"})))
        self.store.db.commit()
        with self.assertRaises(ProviderError): provider.complete([{"role":"user","content":"test"}],max_cost=0.99,budget_scope="RUN-budget")
        self.store.db.execute("INSERT INTO usage(provider,model,timestamp,body) VALUES(?,?,?,?)",("openai","old","test",canonical({"budget_scope":"RUN-unknown","status":"started","cost":None})))
        self.store.db.commit()
        with self.assertRaises(ProviderError): provider.complete([{"role":"user","content":"test"}],max_cost=1,budget_scope="RUN-unknown")

    def test_usage_intent_is_durable_before_request(self):
        def handler(request):
            body=json.loads(self.store.db.execute("SELECT body FROM usage ORDER BY seq DESC LIMIT 1").fetchone()[0])
            self.assertEqual(body["status"],"started"); return self.response()
        self.provider(handler).complete([{"role":"user","content":"test"}],[mock_tool()])
        self.assertEqual(self.store.db.execute("SELECT COUNT(*) FROM usage").fetchone()[0],1)

    def test_endpoint_and_extra_body_guards(self):
        for url in ("http://example.com","https://key@example.com","https://example.com?secret=x"):
            self.config["base_url"]=url
            with self.assertRaises(ValueError): Provider(self.config)
        self.config["base_url"]="https://example.com"; self.config["extra_body"]={"model":"other"}
        with self.assertRaises(ValueError): Provider(self.config)

class KnowledgeCase(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.proof={"SYNTHETIC-REFERENCE":{"mode":"real","passed":True,"evidence_manifest_sha256":"a"*64}}
        self.index=KnowledgeIndex(Path(self.tmp.name)/"index.sqlite",self.proof); self.addCleanup(self.index.close)
        self.source=dict(schema_revision=VERSION,source_id="SRC-TEST",title="Synthetic software reference",source_type="local_observation",url="https://example.invalid/test",locator="test fixture",software_version="5.6",retrieved_at="2026-10-19",access_scope="metadata",review_status="verified",snapshot_sha256=None)
        self.record=dict(schema_revision=VERSION,knowledge_id="KNOW-TEST",revision=1,kind="software_method",claim="MaximumLikelihood 分类 ERR42",source_ids=["SRC-TEST"],source_locator="MaximumLikelihood",software="ENVI",software_version_scope=">=5.6,<5.7",sensor_product_scope=[],required_licenses=["ENVI"],input_formats=["ENVI"],preconditions=[],limitations=["synthetic fixture"],implementation_kind="native",source_review_status="verified",local_verification_status="verified",verification_run_ids=["SYNTHETIC-REFERENCE"],conflicts_with=[],reviewer_role="cloud_codex",lifecycle="published")
        self.index.ingest([self.record],[self.source])

    def search(self,**kwargs):
        return self.index.search("MaximumLikelihood",kwargs.get("software","ENVI"),kwargs.get("version","5.6"),kwargs.get("licenses",["ENVI"]),input_format=kwargs.get("input_format","ENVI"),include_candidates=kwargs.get("include_candidates",False))

    def test_filters_precede_recommendation(self):
        self.assertEqual(len(self.search()),1)
        for kwargs in ({"version":"5.5"},{"version":"unknown"},{"licenses":[]},{"input_format":"GeoTIFF"},{"software":"MATLAB"}):
            with self.subTest(kwargs=kwargs): self.assertEqual(self.search(**kwargs),[])

    def test_candidates_do_not_enter_default_execution(self):
        r=copy.deepcopy(self.record); r.update(knowledge_id="CANDIDATE",lifecycle="candidate",local_verification_status="not_run",verification_run_ids=[])
        self.index.ingest([r],[])
        self.assertEqual(len(self.search()),1)
        candidates=self.search(include_candidates=True)
        self.assertFalse(next(r for r in candidates if r["knowledge_id"]=="CANDIDATE")["execution_eligible"])

    def test_publication_requires_evidence(self):
        r=copy.deepcopy(self.record); r.update(knowledge_id="INVALID",verification_run_ids=[])
        with self.assertRaises(ValueError): self.index.ingest([r],[])

    def test_conflicts_visible_and_execution_blocked(self):
        r=copy.deepcopy(self.record); r.update(revision=2,conflicts_with=["OTHER"])
        self.index.ingest([r],[]); result=self.search()[0]
        self.assertTrue(result["conflict_requires_review"]); self.assertFalse(result["execution_eligible"])

    def test_withdrawal_does_not_fall_back_to_old_revision(self):
        r=copy.deepcopy(self.record); r.update(revision=2,lifecycle="withdrawn")
        self.index.ingest([r],[]); self.assertEqual(self.search(),[])

    def test_offline_snapshot_restore_and_tamper(self):
        path=Path(self.tmp.name)/"snapshot.json"; self.index.snapshot(path)
        other=KnowledgeIndex(Path(self.tmp.name)/"other.sqlite",self.proof); self.addCleanup(other.close)
        other.restore(path)
        self.assertEqual(len(other.search("ERR42","ENVI","5.6",["ENVI"],input_format="ENVI")),1)
        body=json.loads(path.read_text()); body["snapshot"]["knowledge"][0]["claim"]="tampered"; path.write_text(canonical(body))
        with self.assertRaises(ValueError): other.restore(path)

    def test_exact_matlab_release_scope(self):
        self.assertTrue(compatible("exact:R2024a","R2024a")); self.assertFalse(compatible("exact:R2024a","R2024b"))

    def test_heldout_answers_refused(self):
        r=copy.deepcopy(self.record); r.update(knowledge_id="ANSWER",kind="heldout_answer")
        with self.assertRaises(ValueError): self.index.ingest([r],[])

    def test_source_returned_with_status_and_locator(self):
        result=self.search()[0]; self.assertEqual(result["sources"][0]["url"],self.source["url"])
        self.assertEqual(result["local_verification_status"],"verified")

    def test_unreviewed_run_reference_cannot_publish(self):
        r=copy.deepcopy(self.record); r.update(knowledge_id="UNREVIEWED",verification_run_ids=["not-in-catalog"])
        with self.assertRaises(ValueError): self.index.ingest([r],[])

class ExperienceCase(unittest.TestCase):
    def test_lifecycle_independent_review_and_withdrawal(self):
        with tempfile.TemporaryDirectory() as tmp, closing(ExperienceRegistry(Path(tmp)/"experience.sqlite")) as registry:
            evidence={"REAL-1":{"mode":"real","passed":True,"evidence_manifest_sha256":"a"*64}}
            record=dict(schema_revision=VERSION,experience_id="EXP-1",revision=1,phenomenon="test",scope="test-only",evidence_run_ids=["REAL-1"],reproduction_run_ids=[],fix_comparison_run_ids=[],regression_run_ids=[],reviewer_role=None,author_role="cloud_codex",status="candidate",supersedes=None)
            registry.save(record,evidence); self.assertEqual(registry.active(),[])
            for rev,status,key in ((2,"reproduced","reproduction_run_ids"),(3,"fix_verified","fix_comparison_run_ids"),(4,"published","regression_run_ids")):
                record=copy.deepcopy(record); record.update(revision=rev,supersedes=rev-1,status=status); record[key]=["REAL-1"]
                if rev==4:
                    record["reviewer_role"]="cloud_codex"
                    with self.assertRaises(ValueError): registry.save(record,evidence)
                    record["reviewer_role"]="gemini_antigravity"
                registry.save(record,evidence)
            self.assertEqual(len(registry.active()),1)
            record=copy.deepcopy(record); record.update(revision=5,supersedes=4,status="withdrawn")
            registry.save(record,evidence); self.assertEqual(registry.active(),[])

    def test_mock_cannot_promote_experience(self):
        with tempfile.TemporaryDirectory() as tmp, closing(ExperienceRegistry(Path(tmp)/"experience.sqlite")) as registry:
            record=dict(schema_revision=VERSION,experience_id="EXP-1",revision=1,phenomenon="test",scope="test-only",evidence_run_ids=["MOCK"],reproduction_run_ids=[],fix_comparison_run_ids=[],regression_run_ids=[],reviewer_role=None,author_role="cloud_codex",status="candidate",supersedes=None)
            with self.assertRaises(ValueError): registry.save(record,{"MOCK":{"mode":"mock","passed":True,"evidence_manifest_sha256":"a"*64}})

class ContractCase(unittest.TestCase):
    def test_all_schemas_are_valid(self):
        import contracts
        for path in Path(contracts.__file__).parent.glob("*.schema.json"):
            Draft202012Validator.check_schema(json.loads(path.read_text()))

    def test_path_traversal_and_symlink_rejected(self):
        with tempfile.TemporaryDirectory() as root, tempfile.TemporaryDirectory() as outside:
            for value in ("../file","/tmp/file","C:/file","C:\\file","a\\..\\file"):
                with self.subTest(path=value), self.assertRaises(ValueError): within(root,value)
            if os.name!="nt":
                (Path(root)/"link").symlink_to(outside,target_is_directory=True)
                with self.assertRaises(ValueError): within(root,"link/file")

    def test_text_and_docx_extraction_locations(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"要求.txt"; path.write_text("分类\n需要报告",encoding="utf-8")
            self.assertEqual(extract(path)["blocks"][1]["source_locator"],"要求.txt:line 2")
            from docx import Document
            doc=Document(); doc.add_paragraph("实验要求"); table=doc.add_table(rows=1,cols=1); table.cell(0,0).text="报告"
            path=Path(tmp)/"要求.docx"; doc.save(path)
            result=extract(path); self.assertEqual(len(result["blocks"]),2); self.assertTrue(result["gaps"])

if __name__=="__main__": unittest.main()
